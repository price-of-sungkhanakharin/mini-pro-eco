#!/usr/bin/env python3
"""Export FastAPI OpenAPI spec snapshot into CSV and Excel formats."""

import csv
import json
import sys
from pathlib import Path

# Add project root directory to python import path
ROOT_DIR = Path(__file__).resolve().parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.main import app


def export_openapi_snapshot():
    """Extract OpenAPI schema from FastAPI app and write CSV/Excel snapshots."""
    print("Generating OpenAPI schema from FastAPI application...")
    openapi_schema = app.openapi()

    # Save raw openapi.json for reference
    json_path = ROOT_DIR / "openapi.json"
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(openapi_schema, f, indent=2)
    print(f"Saved raw OpenAPI JSON schema to {json_path}")

    rows = []
    paths = openapi_schema.get("paths", {})

    for path_str, path_item in paths.items():
        for method_str, op_data in path_item.items():
            if method_str.lower() not in ["get", "post", "put", "delete", "patch", "options", "head"]:
                continue

            method = method_str.upper()
            tags = ", ".join(op_data.get("tags", ["Untagged"]))
            summary = op_data.get("summary", "")
            description = op_data.get("description", "").replace("\n", " ").strip()
            operation_id = op_data.get("operationId", "")

            responses_dict = op_data.get("responses", {})
            responses_list = []
            for code, resp in responses_dict.items():
                resp_desc = resp.get("description", "")
                responses_list.append(f"{code}: {resp_desc}")
            responses_str = " | ".join(responses_list)

            deprecated = op_data.get("deprecated", False)

            rows.append({
                "Path": path_str,
                "Method": method,
                "Tag": tags,
                "Summary": summary,
                "Description": description,
                "Operation ID": operation_id,
                "Responses": responses_str,
                "Deprecated": deprecated,
            })

    # Sort rows logically by Tag, Path, Method
    rows.sort(key=lambda r: (r["Tag"], r["Path"], r["Method"]))

    # 1. Export CSV
    csv_path = ROOT_DIR / "openapi_snapshot.csv"
    fieldnames = ["Path", "Method", "Tag", "Summary", "Description", "Operation ID", "Responses", "Deprecated"]
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)
    print(f"Successfully exported CSV snapshot to {csv_path}")

    # 2. Export Excel (XLSX)
    excel_path = ROOT_DIR / "openapi_snapshot.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "OpenAPI Endpoints"

    # Styling header
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    ws.append(fieldnames)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        ws.append([row[col] for col in fieldnames])

    # Adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 65)

    wb.save(excel_path)
    print(f"Successfully exported Excel snapshot to {excel_path}")


if __name__ == "__main__":
    export_openapi_snapshot()
